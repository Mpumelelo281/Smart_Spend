import calendar
from decimal import Decimal
from io import BytesIO

from django.db.models import Sum
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from rest_framework import generics, permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.permissions import IsStudent

from .models import Budget, Transaction
from .serializers import (
    BudgetCategoryCreateSerializer,
    BudgetCategorySerializer,
    BudgetCloneSerializer,
    BudgetCreateSerializer,
    BudgetSerializer,
    TransactionSerializer,
)


class BudgetListCreateView(generics.ListCreateAPIView):
    """Rule 9: RBAC enforced here via permission_classes, not by the client
    hiding the "create budget" button — a Student can only ever see/create
    budgets tied to their own StudentProfile (see get_queryset).
    """

    permission_classes = [permissions.IsAuthenticated, IsStudent]

    def get_queryset(self):
        return Budget.objects.filter(profile__user=self.request.user).order_by("-year", "-month")

    def get_serializer_class(self):
        return BudgetCreateSerializer if self.request.method == "POST" else BudgetSerializer

    def get_serializer_context(self):
        context = super().get_serializer_context()
        if self.request.method == "POST":
            context["profile"] = self.request.user.student_profile
        return context

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        budget = serializer.save()
        return Response(BudgetSerializer(budget).data, status=status.HTTP_201_CREATED)


class BudgetDetailView(generics.RetrieveAPIView):
    permission_classes = [permissions.IsAuthenticated, IsStudent]
    serializer_class = BudgetSerializer

    def get_queryset(self):
        return Budget.objects.filter(profile__user=self.request.user)

    def get_object(self):
        return get_object_or_404(self.get_queryset(), pk=self.kwargs["pk"])


class TransactionCreateView(generics.CreateAPIView):
    permission_classes = [permissions.IsAuthenticated, IsStudent]
    serializer_class = TransactionSerializer
    queryset = Transaction.objects.all()


class BudgetCloneView(APIView):
    """Recurring budgets: clone this budget's categories into a new month
    instead of re-typing them — see BudgetCloneSerializer for the
    validation rules.
    """

    permission_classes = [permissions.IsAuthenticated, IsStudent]

    def post(self, request, pk):
        source = get_object_or_404(Budget, pk=pk, profile__user=request.user)
        serializer = BudgetCloneSerializer(
            data=request.data,
            context={"source_budget": source, "profile": request.user.student_profile},
        )
        serializer.is_valid(raise_exception=True)
        budget = serializer.save()
        return Response(BudgetSerializer(budget).data, status=status.HTTP_201_CREATED)


class BudgetHistoryExportView(APIView):
    """An Excel (.xlsx) download of every budget the student has ever
    created, one row per category, so it's usable in a spreadsheet for
    NSFAS reporting or a student's own records — the same data the History
    page already charts, just exported rather than only ever viewable
    in-app. A real workbook (bold header, banded Excel Table, currency
    formatting) rather than plain CSV — CSV has no concept of formatting at
    all, so "make the header bold" is only possible with an actual
    spreadsheet file format.
    """

    permission_classes = [permissions.IsAuthenticated, IsStudent]

    def get(self, request):
        budgets = (
            Budget.objects.filter(profile__user=request.user)
            .prefetch_related("categories__transactions")
            .order_by("year", "month")
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Budget history"

        headers = [
            "Month",
            "Year",
            "Category",
            "Allocated (R)",
            "Spent (R)",
            "Month Total Allocated (R)",
            "Month Total Spent (R)",
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        row_count = 1
        for budget in budgets:
            # Same total the History page's summary/table already shows
            # (see BudgetSerializer.get_total_spent) — repeated on every
            # category row for that month, not just the category's own
            # figure, so "was R1,100 the whole month or just this
            # category?" is answerable directly from the row.
            month_total_spent = Transaction.objects.filter(category__budget=budget).aggregate(
                total=Coalesce(Sum("amount"), Decimal("0"))
            )["total"]
            for category in budget.categories.all():
                spent = category.transactions.aggregate(total=Coalesce(Sum("amount"), Decimal("0")))["total"]
                sheet.append(
                    [
                        budget.month,
                        budget.year,
                        category.name,
                        float(category.allocated_amount),
                        float(spent),
                        float(budget.total_allocated),
                        float(month_total_spent),
                    ]
                )
                row_count += 1

        for column, width in zip("ABCDEFG", (8, 8, 22, 16, 14, 24, 22)):
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2, min_col=4, max_col=7):
            for cell in row:
                cell.number_format = "R #,##0.00"

        # A real Excel Table (not just formatted cells) so Excel/Sheets
        # render it with banded rows and filter dropdowns — "make it a
        # table", not merely bold text above some numbers. Needs at least
        # one data row; an empty history still gets a plain bold header.
        if row_count > 1:
            table = Table(displayName="BudgetHistory", ref=f"A1:G{row_count}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium9", showRowStripes=True, showFirstColumn=False
            )
            sheet.add_table(table)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="smartspend-budget-history.xlsx"'
        return response


class BudgetForecastView(APIView):
    """Projects each category's (and the budget's) end-of-month spend from
    the student's pace so far this month — not a prediction model, just
    `spent_so_far / days_elapsed * days_in_month`, which is enough to warn
    "you're on track to go over" before it actually happens, complementing
    Rule 10's after-the-fact 80%/exceeded alerts.
    """

    permission_classes = [permissions.IsAuthenticated, IsStudent]

    def get(self, request, pk):
        budget = get_object_or_404(Budget, pk=pk, profile__user=request.user)
        today = timezone.localdate()
        days_in_month = calendar.monthrange(budget.year, budget.month)[1]

        if (budget.year, budget.month) == (today.year, today.month):
            days_elapsed = today.day
        elif (budget.year, budget.month) < (today.year, today.month):
            days_elapsed = days_in_month  # a past month: already complete, no need to project
        else:
            days_elapsed = 0  # a future budget: no spend yet to pace from

        categories = []
        total_spent = Decimal("0")
        total_projected = Decimal("0")
        for category in budget.categories.all():
            spent = category.transactions.aggregate(total=Coalesce(Sum("amount"), Decimal("0")))["total"]
            if days_elapsed > 0:
                projected = (spent / days_elapsed * days_in_month).quantize(Decimal("0.01"))
            else:
                projected = Decimal("0.00")

            total_spent += spent
            total_projected += projected
            categories.append(
                {
                    "category_id": category.category_id,
                    "name": category.name,
                    "allocated_amount": category.allocated_amount,
                    "spent_amount": spent,
                    "projected_amount": projected,
                    "will_exceed": projected > category.allocated_amount,
                }
            )

        return Response(
            {
                "budget_id": budget.budget_id,
                "days_elapsed": days_elapsed,
                "days_in_month": days_in_month,
                "total_allocated": budget.total_allocated,
                "total_spent": total_spent,
                "total_projected": total_projected,
                "will_exceed_allowance": total_projected > budget.total_allocated,
                "categories": categories,
            }
        )


class BudgetCategoryCreateView(APIView):
    """Backs the "Other" option in the log-expense form: lets a student add
    a category to their own existing budget on the fly, for spend that
    doesn't fit anything they pre-allocated. See
    BudgetCategoryCreateSerializer for why this never touches Rule 1.
    """

    permission_classes = [permissions.IsAuthenticated, IsStudent]

    def post(self, request, budget_id):
        budget = get_object_or_404(Budget, pk=budget_id, profile__user=request.user)
        serializer = BudgetCategoryCreateSerializer(data=request.data, context={"budget": budget})
        serializer.is_valid(raise_exception=True)
        category = serializer.save()
        return Response(BudgetCategorySerializer(category).data, status=status.HTTP_201_CREATED)
