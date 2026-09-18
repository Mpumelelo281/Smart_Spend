import uuid
from decimal import Decimal
from io import BytesIO

import pytest
from freezegun import freeze_time
from openpyxl import load_workbook

from apps.budgets.models import Budget

pytestmark = pytest.mark.django_db


def _create_budget(client, month=3, year=2026, allocated="500.00"):
    response = client.post(
        "/api/v1/budgets/",
        {"month": month, "year": year, "categories": [{"name": "Food", "allocated_amount": allocated}]},
        format="json",
    )
    return response.data["budget_id"]


# --- clone -----------------------------------------------------------------


def test_clone_copies_categories_into_new_month(authenticated_client):
    budget_id = _create_budget(authenticated_client)

    response = authenticated_client.post(f"/api/v1/budgets/{budget_id}/clone/", {"month": 4, "year": 2026})
    assert response.status_code == 201
    assert response.data["month"] == 4
    assert response.data["categories"][0]["name"] == "Food"
    assert str(response.data["categories"][0]["allocated_amount"]) == "500.00"


def test_clone_rejects_existing_target_month(authenticated_client):
    budget_id = _create_budget(authenticated_client)
    _create_budget(authenticated_client, month=4)

    response = authenticated_client.post(f"/api/v1/budgets/{budget_id}/clone/", {"month": 4, "year": 2026})
    assert response.status_code == 400


def test_clone_rejects_when_exceeding_current_allowance(authenticated_client, student_user):
    budget_id = _create_budget(authenticated_client, allocated="1600.00")

    profile = student_user.student_profile
    profile.allowance_amount = "1000.00"
    profile.save()

    response = authenticated_client.post(f"/api/v1/budgets/{budget_id}/clone/", {"month": 4, "year": 2026})
    assert response.status_code == 400


def test_clone_only_own_budget(authenticated_client):
    response = authenticated_client.post(
        f"/api/v1/budgets/{uuid.uuid4()}/clone/", {"month": 4, "year": 2026}
    )
    assert response.status_code == 404


# --- export ------------------------------------------------------------


def test_export_returns_an_excel_table_with_a_bold_header(authenticated_client):
    _create_budget(authenticated_client, month=3, allocated="500.00")

    response = authenticated_client.get("/api/v1/budgets/export/")
    assert response.status_code == 200
    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response["Content-Disposition"] == 'attachment; filename="smartspend-budget-history.xlsx"'

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active

    header = [cell.value for cell in sheet[1]]
    assert header == [
        "Month",
        "Year",
        "Category",
        "Allocated (R)",
        "Spent (R)",
        "Month Total Allocated (R)",
        "Month Total Spent (R)",
    ]
    assert all(cell.font.bold for cell in sheet[1])

    data_row = [cell.value for cell in sheet[2]]
    assert data_row == [3, 2026, "Food", 500.0, 0.0, 500.0, 0.0]

    assert "BudgetHistory" in sheet.tables
    assert sheet.tables["BudgetHistory"].ref == "A1:G2"


def test_export_month_total_columns_sum_across_categories(authenticated_client):
    # A month with two categories: the per-category Spent column should
    # differ per row, but Month Total Spent must be the same on every row
    # for that month — exactly the "was it just this category, or the
    # whole month?" ambiguity this column exists to remove.
    budget_id = _create_budget(authenticated_client, month=9, allocated="350.00")
    authenticated_client.post(
        f"/api/v1/budgets/{budget_id}/categories/", {"name": "Toiletries"}, format="json"
    )
    toiletries = Budget.objects.get(pk=budget_id).categories.get(name="Toiletries")

    authenticated_client.post(
        "/api/v1/budgets/transactions/",
        {"category": toiletries.category_id, "amount": "1100.00", "purchase_date": "2026-09-05"},
        format="json",
    )

    response = authenticated_client.get("/api/v1/budgets/export/")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active

    rows = {row[2]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows["Food"][4] == 0.0  # this category's own spend
    assert rows["Toiletries"][4] == 1100.0  # this category's own spend
    assert rows["Food"][6] == 1100.0  # whole month's total, same on both rows
    assert rows["Toiletries"][6] == 1100.0


def test_export_with_no_budgets_still_has_a_bold_header(authenticated_client):
    response = authenticated_client.get("/api/v1/budgets/export/")
    assert response.status_code == 200

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.max_row == 1


def test_export_requires_authentication(api_client):
    response = api_client.get("/api/v1/budgets/export/")
    assert response.status_code == 401


# --- forecast ------------------------------------------------------------


def _authenticate(api_client, user):
    # Minted *inside* the frozen-time block (unlike the authenticated_client
    # fixture, which resolves before the test body runs) so the token's
    # `iat` lines up with the frozen "now" — otherwise a freeze_time date
    # earlier than the real wall-clock moment the fixture ran makes the
    # token look issued in the future and JWT auth rejects it.
    from rest_framework_simplejwt.tokens import RefreshToken

    token = RefreshToken.for_user(user)
    api_client.credentials(HTTP_AUTHORIZATION=f"Bearer {token.access_token}")
    return api_client


def test_forecast_projects_spend_from_pace_this_month(api_client, student_user):
    with freeze_time("2026-03-10"):
        client = _authenticate(api_client, student_user)
        budget_id = _create_budget(client, month=3, allocated="300.00")
        category_id = Budget.objects.get(pk=budget_id).categories.first().category_id
        client.post(
            "/api/v1/budgets/transactions/",
            {"category": category_id, "amount": "50.00", "purchase_date": "2026-03-05"},
            format="json",
        )

        response = client.get(f"/api/v1/budgets/{budget_id}/forecast/")
        assert response.status_code == 200
        assert response.data["days_elapsed"] == 10
        # 50 spent over 10 days in a 31-day month => 155.00 projected.
        assert response.data["total_projected"] == Decimal("155.00")
        assert response.data["will_exceed_allowance"] is False


def test_forecast_flags_categories_on_track_to_exceed(api_client, student_user):
    with freeze_time("2026-03-10"):
        client = _authenticate(api_client, student_user)
        budget_id = _create_budget(client, month=3, allocated="100.00")
        category_id = Budget.objects.get(pk=budget_id).categories.first().category_id
        client.post(
            "/api/v1/budgets/transactions/",
            {"category": category_id, "amount": "80.00", "purchase_date": "2026-03-05"},
            format="json",
        )

        response = client.get(f"/api/v1/budgets/{budget_id}/forecast/")
        assert response.data["categories"][0]["will_exceed"] is True


def test_forecast_for_a_completed_past_month_uses_full_month(api_client, student_user):
    with freeze_time("2026-04-15"):
        client = _authenticate(api_client, student_user)
        budget_id = _create_budget(client, month=3, allocated="300.00")

        response = client.get(f"/api/v1/budgets/{budget_id}/forecast/")
        assert response.data["days_elapsed"] == 31  # March has 31 days, already finished
